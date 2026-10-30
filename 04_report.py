from pathlib import Path

import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill

YM = "202610"
YEAR, MONTH = int(YM[:4]), int(YM[4:])

# ① 3-Bで作ったきれいなデータを読む
read_opts = dict(encoding="utf-8-sig", dtype={"商品コード": str}, parse_dates=["注文日"])
sales = pd.read_csv(f"output/sales_{YM}.csv", **read_opts)
returns = pd.read_csv(f"output/returns_{YM}.csv", **read_opts)
log = pd.read_csv(f"output/log_{YM}.csv", encoding="utf-8-sig")

# ② 集計する
# 送料は「注文ごとに1回だけ」数える（1注文に複数行あっても同じ送料が並んでいるため）
shipping = sales.groupby("注文番号")["送料"].first().sum()

summary = pd.DataFrame(
    {
        "項目": ["売上合計（税込・送料除く）", "注文件数", "販売個数", "送料の合計", "返品（件数）", "返品（金額）"],
        "値": [
            sales["小計"].sum(),
            sales["注文番号"].nunique(),
            sales["数量"].sum(),
            shipping,
            len(returns),
            returns["小計"].sum(),
        ],
    }
)

by_product = (
    sales.groupby(["商品コード", "商品名"], as_index=False)
    .agg(販売個数=("数量", "sum"), 売上=("小計", "sum"))
    .sort_values("売上", ascending=False)
)

by_category = (
    sales.groupby("カテゴリ", as_index=False)
    .agg(販売個数=("数量", "sum"), 売上=("小計", "sum"))
    .sort_values("売上", ascending=False)
)

# 日別：注文がなかった日も「0円」として並べる
days = pd.date_range(f"{YEAR}-{MONTH:02d}-01", periods=pd.Period(YM, "M").days_in_month)
by_day = sales.groupby("注文日")["小計"].sum().reindex(days, fill_value=0)
by_day = by_day.rename_axis("日付").reset_index(name="売上")
by_day["日付"] = by_day["日付"].dt.strftime("%m/%d")

returns_list = returns[["注文番号", "注文日", "商品名", "数量", "小計"]].copy()
returns_list["注文日"] = returns_list["注文日"].dt.strftime("%Y/%m/%d")

# ③ Excelに書き出す（シートごとに表を置く）
Path("report").mkdir(exist_ok=True)
out = f"report/売上レポート_{YEAR}年{MONTH}月.xlsx"
sheets = {
    "サマリー": summary,
    "商品別": by_product,
    "カテゴリ別": by_category,
    "日別": by_day,
    "返品": returns_list,
    "修正ログ": log,
}

with pd.ExcelWriter(out, engine="openpyxl") as writer:
    for name, table in sheets.items():
        table.to_excel(writer, sheet_name=name, index=False)

    # ④ 見た目を整える
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for name, ws in writer.sheets.items():
        for cell in ws[1]:  # 1行目（見出し）
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for col in ws.columns:  # 列の幅を中身に合わせる
            width = max(len(str(c.value)) for c in col if c.value is not None)
            ws.column_dimensions[col[0].column_letter].width = min(width * 2 + 2, 60)
        for row in ws.iter_rows(min_row=2):  # 数字に3桁カンマ
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0"
        ws.freeze_panes = "A2"  # 見出しを固定

    # ⑤ 日別シートに折れ線グラフを付ける
    ws = writer.sheets["日別"]
    chart = LineChart()
    chart.title = f"{MONTH}月の日別売上"
    chart.y_axis.title = "売上（円）"
    data = Reference(ws, min_col=2, min_row=1, max_row=len(by_day) + 1)
    labels = Reference(ws, min_col=1, min_row=2, max_row=len(by_day) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    chart.series[0].smooth = False  # 線をなめらかにしない（実際の値どおりに結ぶ）
    chart.legend = None

    chart.varyColors = False  # 線を1色にする（点ごとに色を変えない）
    chart.series[0].graphicalProperties.line.solidFill = "4472C4"  # 線の色（青）
    chart.x_axis.delete = False  # 横軸（日付）の目盛りを表示する
    chart.y_axis.delete = False  # 縦軸（金額）の目盛りを表示する
    chart.y_axis.numFmt = "#,##0"
    chart.width, chart.height = 24, 10
    ws.add_chart(chart, "D2")

print(f"✅ レポートを作成しました：{out}")
print()
print(summary.to_string(index=False))